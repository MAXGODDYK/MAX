from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent
OUT = BASE_DIR / "Точные рецепты табачных смесей для кальяна.docx"


RECIPES = [
    {
        "name": "Лёгкая дымная",
        "profile": "низкая крепость, высокая дымность, средняя жаростойкость",
        "use": "мягкий тестовый профиль, когда важнее дым и лёгкость, чем плотная табачная крепость",
        "ingredients": [
            ("Табак", 22, 220, "сушеный и нарезанный; мягкая Virginia или аналогичная лёгкая база"),
            ("Глицерин", 36, 360, "пищевой/фармацевтический; основной вклад в дымность"),
            ("ГФС", 26, 260, "глюкозно-фруктозный сироп; сладость, вязкость, удержание влаги"),
            ("Пропиленгликоль", 6, 60, "раскрытие аромата и равномерность пропитки"),
            ("Дистиллированная вода", 3, 30, "малая корректировка вязкости; не делать основной увлажнитель"),
            ("Ароматизатор", 7, 70, "пищевой водорастворимый; начинать ниже при сильном концентрате"),
        ],
        "sous_vide": "50 °C, 2 часа в герметичном пакете или банке; затем 24-48 часов созревания.",
        "multi": "Режим подогрева/йогурта с водяной баней 45-55 °C, 1.5-2 часа; контролировать термометром.",
        "notes": "Если смесь кажется слишком жидкой, дай ей отстояться открытой 20-40 минут и перемешай. Воду выше 3% не поднимать без малого теста.",
    },
    {
        "name": "Средняя сбалансированная",
        "profile": "средняя крепость, высокая дымность, жаростойкость выше средней",
        "use": "универсальная формула для сравнения табака, сиропа и ароматизаторов",
        "ingredients": [
            ("Табак", 30, 300, "сушеный и нарезанный; Virginia/Burley blend или средняя крепость"),
            ("Глицерин", 32, 320, "пищевой/фармацевтический; дымность без чрезмерной водянистости"),
            ("ГФС", 24, 240, "глюкозно-фруктозный сироп; вязкость и сладкая основа"),
            ("Пропиленгликоль", 5, 50, "растворение и перенос аромата"),
            ("Дистиллированная вода", 3, 30, "малая корректировка вязкости жидкой основы"),
            ("Ароматизатор", 6, 60, "пищевой водорастворимый; универсальная стартовая дозировка"),
        ],
        "sous_vide": "52 °C, 2.5 часа; затем 48 часов созревания с перемешиванием через сутки.",
        "multi": "Водяная баня 48-55 °C, 2 часа; банку не перегревать и не ставить прямо на горячее дно.",
        "notes": "Оптимальная формула для первой партии на 100-250 г перед масштабированием. Если лист сухой, лучше дольше созревать, а не сразу добавлять много воды.",
    },
    {
        "name": "Крепкая жаростойкая",
        "profile": "высокая крепость, средняя дымность, высокая жаростойкость",
        "use": "плотная смесь для аккуратного прогрева и более выраженного табачного профиля",
        "ingredients": [
            ("Табак", 38, 380, "сушеный и нарезанный; Burley/тёмная Virginia или крепкая база"),
            ("Глицерин", 28, 280, "пищевой/фармацевтический; достаточно дыма без лишней текучести"),
            ("ГФС", 23, 230, "глюкозно-фруктозный сироп; плотность и устойчивость к жару"),
            ("Пропиленгликоль", 3, 30, "умеренное раскрытие аромата"),
            ("Дистиллированная вода", 3, 30, "малая корректировка вязкости; не разжижать плотную сборку"),
            ("Ароматизатор", 5, 50, "пищевой водорастворимый; не перегружать вкусом"),
        ],
        "sous_vide": "55 °C, 2.5-3 часа; затем 48-72 часа созревания.",
        "multi": "Водяная баня 50-55 °C, 2-2.5 часа; обязательно перемешать после прогрева.",
        "notes": "Тестировать с меньшим жаром. Если вкус резкий, увеличить время созревания.",
    },
    {
        "name": "Лёгкая Virginia Gold",
        "profile": "низкая крепость, очень высокая дымность, средняя жаростойкость",
        "use": "светлая мягкая сборка под Virginia Gold и лёгкие ароматизаторы",
        "ingredients": [
            ("Virginia Gold", 20, 200, "лёгкая светлая табачная база; товар выбрать в калькуляторе"),
            ("Глицерин", 38, 380, "максимальная дымность без сильного табачного удара"),
            ("ГФС", 27, 270, "сладкая вязкая часть для удержания влаги"),
            ("Пропиленгликоль", 7, 70, "раскрытие лёгких ароматов"),
            ("Дистиллированная вода", 3, 30, "корректировка текучести жидкой основы"),
            ("Ароматизатор", 5, 50, "мягкая стартовая дозировка"),
        ],
        "sous_vide": "50 °C, 2 часа; затем 24-48 часов созревания.",
        "multi": "Водяная баня 45-52 °C, 1.5-2 часа; перегрев не нужен.",
        "notes": "Если вкус слишком лёгкий, усиливать табак лучше заменой сорта, а не ростом ароматизатора.",
    },
    {
        "name": "Мягкая Virginia дымная",
        "profile": "крепость ниже средней, очень высокая дымность, средняя жаростойкость",
        "use": "универсальная дымная сборка для Virginia средней крепости",
        "ingredients": [
            ("Virginia", 24, 240, "средняя Virginia; товар выбрать отдельно в калькуляторе"),
            ("Глицерин", 36, 360, "плотный дым и мягкость"),
            ("ГФС", 26, 260, "вязкость и сладкая основа"),
            ("Пропиленгликоль", 6, 60, "перенос аромата"),
            ("Дистиллированная вода", 3, 30, "малая корректировка вязкости"),
            ("Ароматизатор", 5, 50, "не перегружать вкус листа"),
        ],
        "sous_vide": "52 °C, 2 часа; затем 48 часов созревания.",
        "multi": "Водяная баня 48-53 °C, 2 часа; перемешать после прогрева.",
        "notes": "Хорошая точка сравнения разных Virginia: меняй только товар табака, не формулу.",
    },
    {
        "name": "Баланс Virginia + Burley",
        "profile": "средняя крепость, высокая дымность, жаростойкость выше средней",
        "use": "сбалансированный микс: мягкость Virginia плюс плотность Burley",
        "ingredients": [
            ("Virginia", 22, 220, "мягкая часть табачной базы; выбрать отдельный товар"),
            ("Burley", 10, 100, "усиление крепости и плотности; выбрать отдельный товар"),
            ("Глицерин", 31, 310, "дымность без лишней текучести"),
            ("ГФС", 24, 240, "вязкость и сладкая основа"),
            ("Пропиленгликоль", 5, 50, "перенос аромата"),
            ("Дистиллированная вода", 3, 30, "малая корректировка вязкости"),
            ("Ароматизатор", 5, 50, "универсальная дозировка"),
        ],
        "sous_vide": "52 °C, 2.5 часа; затем 48 часов созревания.",
        "multi": "Водяная баня 48-55 °C, 2 часа; перемешать и дать отстояться сутки.",
        "notes": "В калькуляторе выбрать два разных товара табака: отдельно Virginia и отдельно Burley.",
    },
    {
        "name": "Крепкая Burley + Virginia",
        "profile": "высокая крепость, средняя дымность, высокая жаростойкость",
        "use": "выраженный табачный профиль с Burley как главным листом",
        "ingredients": [
            ("Burley", 24, 240, "основная крепкая часть; выбрать отдельный товар"),
            ("Virginia", 14, 140, "смягчение профиля и баланс вкуса"),
            ("Глицерин", 28, 280, "умеренная дымность для плотной сборки"),
            ("ГФС", 23, 230, "плотная сиропная часть"),
            ("Пропиленгликоль", 3, 30, "минимально для раскрытия аромата"),
            ("Дистиллированная вода", 3, 30, "малая корректировка вязкости"),
            ("Ароматизатор", 5, 50, "не перебивать табачную базу"),
        ],
        "sous_vide": "55 °C, 2.5-3 часа; затем 48-72 часа созревания.",
        "multi": "Водяная баня 50-55 °C, 2-2.5 часа; избегать перегрева.",
        "notes": "Начинать тест с меньшего жара. Burley обычно требует больше времени на созревание.",
    },
    {
        "name": "Жаростойкая Burley",
        "profile": "высокая крепость, средняя дымность, очень высокая жаростойкость",
        "use": "самая плотная сборка под Burley и аккуратный долгий прогрев",
        "ingredients": [
            ("Burley", 40, 400, "крепкая плотная табачная база; товар выбрать в калькуляторе"),
            ("Глицерин", 27, 270, "умеренная дымность без лишней жидкости"),
            ("ГФС", 23, 230, "плотность и жаростойкость"),
            ("Пропиленгликоль", 3, 30, "минимальное раскрытие аромата"),
            ("Дистиллированная вода", 3, 30, "малая корректировка вязкости"),
            ("Ароматизатор", 4, 40, "умеренно, чтобы не получить резкий вкус"),
        ],
        "sous_vide": "55 °C, 3 часа; затем 72 часа созревания с перемешиванием через сутки.",
        "multi": "Водяная баня 50-55 °C, 2.5 часа; прогревать без кипения.",
        "notes": "Плотная формула: если смесь кажется сухой, сначала увеличить созревание, а не воду.",
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table(table, header_fill: str = "D9EAD3") -> None:
    table.style = "Table Grid"
    table.autofit = True
    for cell in table.rows[0].cells:
        set_cell_shading(cell, header_fill)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def add_note(document: Document, title: str, text: str, fill: str = "FFF2CC") -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title + ". ")
    run.bold = True
    p.add_run(text)
    document.add_paragraph()


def add_recipe(document: Document, recipe: dict) -> None:
    document.add_heading(recipe["name"], level=1)
    add_note(document, "Профиль", recipe["profile"] + ". Назначение: " + recipe["use"] + ".", "EAF3F8")

    table = document.add_table(rows=1, cols=4)
    headers = ["Ингредиент", "Процент", "На 1000 г", "Комментарий"]
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True)
    for ingredient, percent, grams, comment in recipe["ingredients"]:
        cells = table.add_row().cells
        set_cell_text(cells[0], ingredient)
        set_cell_text(cells[1], f"{percent}%")
        set_cell_text(cells[2], f"{grams} г")
        set_cell_text(cells[3], comment)
    total = table.add_row().cells
    set_cell_text(total[0], "Итого", bold=True)
    set_cell_text(total[1], "100%", bold=True)
    set_cell_text(total[2], "1000 г", bold=True)
    set_cell_text(total[3], "Масштабируется пропорционально", bold=True)
    for cell in total:
        set_cell_shading(cell, "F3F6FA")
    style_table(table)

    document.add_heading("Технология", level=2)
    steps = [
        "Взвесить все ингредиенты. Для первого теста лучше сделать 100-250 г, а не сразу большой объём.",
        "Смешать глицерин, ГФС, пропиленгликоль, дистиллированную воду и ароматизатор до однородной жидкой основы.",
        "Добавить основу к табаку небольшими порциями, тщательно перемешивая после каждой порции.",
        "Прогреть одним из способов ниже, не допуская кипения и резкого перегрева.",
        "После прогрева остудить, закрыть и оставить на созревание. Перед тестом перемешать.",
    ]
    for index, step in enumerate(steps, start=1):
        document.add_paragraph(f"{index}. {step}")

    tech = document.add_table(rows=1, cols=2)
    set_cell_text(tech.rows[0].cells[0], "Сувид", bold=True)
    set_cell_text(tech.rows[0].cells[1], "Мультиварка", bold=True)
    cells = tech.add_row().cells
    set_cell_text(cells[0], recipe["sous_vide"])
    set_cell_text(cells[1], recipe["multi"])
    style_table(tech, "DDEBF7")

    document.add_paragraph()
    add_note(document, "Корректировка", recipe["notes"], "FCE4D6")


def build() -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Heading 1", "Heading 2", "Title"]:
        styles[style_name].font.name = "Arial"
    styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
    styles["Heading 2"].font.color.rgb = RGBColor(55, 86, 35)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Точные рецепты табачных смесей для кальяна")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(31, 78, 121)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Домашние техкарты для сушёного и нарезанного табака: проценты + пример на 1000 г").italic = True

    document.add_paragraph()
    add_note(
        document,
        "Важно",
        "Табачные продукты вредят здоровью и могут быть незаконны для несовершеннолетних. Эти материалы предназначены только для взрослых пользователей там, где приготовление и использование таких смесей законно. Документ не является медицинской, юридической или производственной рекомендацией.",
        "F4CCCC",
    )

    document.add_heading("Безопасность и ограничения", level=1)
    safety_points = [
        "Работать только в чистой зоне, с отдельной посудой, перчатками и хорошей вентиляцией.",
        "Использовать ингредиенты пищевого или фармацевтического качества с понятным составом и назначением.",
        "Не использовать масляные ароматизаторы: для кальянных смесей нужны водорастворимые пищевые ароматизаторы.",
        "Не перегревать смесь. Практический диапазон для домашнего прогрева: 45-55 °C; кипение и жарка недопустимы.",
        "Не хранить смесь в открытой таре. Подходят чистые стеклянные банки или пищевые контейнеры с маркировкой даты.",
        "Любую новую формулу сначала проверять на малой партии 100-250 г и вести записи по вкусу, влажности и жаростойкости.",
    ]
    for point in safety_points:
        document.add_paragraph(point, style="List Bullet")

    document.add_heading("Базовая логика масштабирования", level=1)
    document.add_paragraph(
        "Все рецепты ниже заданы в процентах. Пример на 1000 г нужен как удобная техкарта: если нужен другой объём, умножай общий вес на процент ингредиента и дели на 100. Например, 32% от 250 г = 80 г."
    )

    quick = document.add_table(rows=1, cols=4)
    for idx, header in enumerate(["Цель", "Больше", "Меньше", "Что проверить"]):
        set_cell_text(quick.rows[0].cells[idx], header, bold=True)
    quick_rows = [
        ("Дымность", "глицерин", "табак/сироп", "не стала ли смесь слишком жидкой"),
        ("Крепость", "табак", "глицерин/сироп", "не стал ли вкус резким"),
        ("Жаростойкость", "табак и плотная сиропная часть", "лишний PG/ароматизатор", "не горит ли верхний слой"),
        ("Текучесть", "вода по 1-2 г на тест", "вода/PG", "нет ли стекания сиропа в чашу"),
    ]
    for row in quick_rows:
        cells = quick.add_row().cells
        for idx, text in enumerate(row):
            set_cell_text(cells[idx], text)
    style_table(quick, "E2F0D9")

    for recipe in RECIPES:
        document.add_page_break()
        add_recipe(document, recipe)

    document.add_page_break()
    document.add_heading("Журнал тестов", level=1)
    journal = document.add_table(rows=1, cols=6)
    for idx, header in enumerate(["Дата", "Рецепт", "Партия, г", "Влажность", "Жар", "Заметки"]):
        set_cell_text(journal.rows[0].cells[idx], header, bold=True)
    for _ in range(8):
        cells = journal.add_row().cells
        for cell in cells:
            set_cell_text(cell, "")
    style_table(journal, "D9EAD3")

    document.core_properties.title = "Точные рецепты табачных смесей для кальяна"
    document.core_properties.subject = "Техкарты, безопасность, масштабирование"
    document.core_properties.author = "Codex"
    document.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
